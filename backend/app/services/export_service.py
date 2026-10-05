import io
import re
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


class ExportService:
    @staticmethod
    def html_to_docx(html_content: str, title: str = "Document") -> io.BytesIO:
        doc = Document()
        
        # Page Margins: 1 inch
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            
        soup = BeautifulSoup(html_content or "", "html.parser")
        
        # Document Title
        doc.add_heading(title or "Document", level=0)
        
        for element in soup.find_all(['h1', 'h2', 'h3', 'p', 'ul', 'ol']):
            text = element.get_text().strip()
            if not text:
                continue
            if element.name == 'h1':
                doc.add_heading(text, level=1)
            elif element.name == 'h2':
                doc.add_heading(text, level=2)
            elif element.name == 'h3':
                doc.add_heading(text, level=3)
            elif element.name == 'p':
                p = doc.add_paragraph(text)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(6)
            elif element.name in ['ul', 'ol']:
                for li in element.find_all('li'):
                    li_text = li.get_text().strip()
                    if li_text:
                        doc.add_paragraph(li_text, style='List Bullet')
                        
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def html_to_pdf(html_content: str, title: str = "Document") -> io.BytesIO:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=72)
        styles = getSampleStyleSheet()
        
        story = []
        soup = BeautifulSoup(html_content or "", "html.parser")
        
        title_style = styles["Title"]
        h1_style = styles["Heading1"]
        h2_style = styles["Heading2"]
        body_style = styles["Normal"]
        body_style.fontSize = 11
        body_style.leading = 14
        
        story.append(Paragraph(title or "Document", title_style))
        story.append(Spacer(1, 14))
        
        for element in soup.find_all(['h1', 'h2', 'h3', 'p', 'ul', 'ol']):
            text = re.sub(r'\s+', ' ', element.get_text().strip())
            if not text:
                continue
            if element.name in ['h1', 'h2']:
                story.append(Spacer(1, 10))
                story.append(Paragraph(text, h1_style))
                story.append(Spacer(1, 6))
            elif element.name == 'h3':
                story.append(Spacer(1, 8))
                story.append(Paragraph(text, h2_style))
                story.append(Spacer(1, 4))
            elif element.name == 'p':
                story.append(Paragraph(text, body_style))
                story.append(Spacer(1, 6))
            elif element.name in ['ul', 'ol']:
                for li in element.find_all('li'):
                    li_text = li.get_text().strip()
                    if li_text:
                        story.append(Paragraph(f"• {li_text}", body_style))
                        story.append(Spacer(1, 3))
                        
        doc.build(story)
        buffer.seek(0)
        return buffer


export_service = ExportService()
